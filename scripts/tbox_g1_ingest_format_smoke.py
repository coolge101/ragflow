#!/usr/bin/env python3
"""G1 多格式入库冒烟：上传 → 解析 → 检索（供 docs/TBOX_INGEST_FORMAT_SMOKE.md 填表）。"""

from __future__ import annotations

import json
import os
import sys
import time
from dataclasses import asdict, dataclass
from pathlib import Path

import requests
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from api.utils.crypt import crypt  # noqa: E402

BASE = os.environ.get("TBOX_SMOKE_BASE_URL", "http://127.0.0.1:9380").rstrip("/")
EMAIL = os.environ.get("TBOX_SMOKE_EMAIL", "admin@ragflow.io")
PASSWORD = os.environ.get("TBOX_SMOKE_PASSWORD", "admin")
DATE_TAG = os.environ.get("TBOX_SMOKE_DATE", "20260524")
PARSE_TIMEOUT_SEC = int(os.environ.get("TBOX_SMOKE_PARSE_TIMEOUT", "180"))


@dataclass
class FormatResult:
    format: str
    keyword: str
    upload: bool
    parse: bool
    search_hit: bool
    run_status: str
    chunk_count: int
    note: str


def login() -> str:
    r = requests.post(
        f"{BASE}/api/v1/auth/login",
        json={"email": EMAIL, "password": crypt(PASSWORD)},
        timeout=30,
    )
    body = r.json()
    if body.get("code") != 0:
        raise RuntimeError(f"login failed: {body.get('message')}")
    auth = r.headers.get("Authorization")
    if not auth:
        raise RuntimeError("login missing Authorization header")
    return auth


def headers(auth: str) -> dict[str, str]:
    return {"Authorization": auth, "Content-Type": "application/json"}


def create_dataset(auth: str, name: str) -> str:
    r = requests.post(
        f"{BASE}/api/v1/datasets",
        headers=headers(auth),
        json={"name": name, "permission": "me", "chunk_method": "naive"},
        timeout=30,
    )
    body = r.json()
    if body.get("code") != 0:
        raise RuntimeError(f"create dataset failed: {body.get('message')}")
    ds_id = body.get("data", {}).get("id")
    if not ds_id:
        raise RuntimeError("create dataset missing id")
    return str(ds_id)


def make_samples(out_dir: Path) -> dict[str, tuple[Path, str]]:
    out_dir.mkdir(parents=True, exist_ok=True)
    specs: dict[str, tuple[str, str]] = {
        "PDF": ("smoke.pdf", f"TBOX-SMOKE-{DATE_TAG}-PDF"),
        "Word": ("smoke.docx", f"TBOX-SMOKE-{DATE_TAG}-DOCX"),
        "Excel": ("smoke.xlsx", f"TBOX-SMOKE-{DATE_TAG}-XLSX"),
        "图片": ("smoke.png", f"TBOX-SMOKE-{DATE_TAG}-PNG"),
    }
    out: dict[str, tuple[Path, str]] = {}

    for label, (fname, kw) in specs.items():
        path = out_dir / fname
        if label == "PDF":
            c = canvas.Canvas(str(path))
            c.drawString(72, 720, f"G1 smoke sample. Keyword: {kw}")
            c.save()
        elif label == "Word":
            doc = Document()
            doc.add_paragraph(f"G1 smoke Word sample. Keyword: {kw}")
            doc.save(path)
        elif label == "Excel":
            wb = Workbook()
            ws = wb.active
            # ExcelParser treats row 1 as header and row 2+ as data rows.
            ws["A1"] = "Keyword"
            ws["A2"] = kw
            wb.save(path)
        else:
            img = Image.new("RGB", (480, 120), color=(240, 248, 255))
            draw = ImageDraw.Draw(img)
            draw.text((12, 40), kw, fill=(15, 23, 42))
            img.save(path)
        out[label] = (path, kw)
    return out


def upload_file(auth: str, dataset_id: str, path: Path) -> str | None:
    with path.open("rb") as f:
        r = requests.post(
            f"{BASE}/api/v1/datasets/{dataset_id}/documents",
            headers={"Authorization": auth},
            files={"file": (path.name, f)},
            timeout=120,
        )
    body = r.json()
    if body.get("code") != 0:
        return None
    data = body.get("data")
    if isinstance(data, list) and data:
        return str(data[0].get("id") or "")
    if isinstance(data, dict):
        return str(data.get("id") or "")
    return None


def trigger_parse(auth: str, dataset_id: str, doc_id: str) -> bool:
    r = requests.post(
        f"{BASE}/api/v1/datasets/{dataset_id}/documents/parse",
        headers=headers(auth),
        json={"document_ids": [doc_id]},
        timeout=30,
    )
    return r.json().get("code") == 0


def wait_parse(auth: str, dataset_id: str, doc_id: str) -> tuple[str, int]:
    deadline = time.time() + PARSE_TIMEOUT_SEC
    while time.time() < deadline:
        r = requests.get(
            f"{BASE}/api/v1/datasets/{dataset_id}/documents",
            headers={"Authorization": auth},
            params={"page": 1, "page_size": 50},
            timeout=30,
        )
        body = r.json()
        if body.get("code") != 0:
            time.sleep(3)
            continue
        for doc in body.get("data", {}).get("docs") or []:
            if str(doc.get("id")) == doc_id:
                run = str(doc.get("run") or "")
                chunks = int(doc.get("chunk_count") or 0)
                if run in ("3", "DONE"):
                    return run, chunks
                if run in ("4", "FAIL"):
                    return run, chunks
        time.sleep(3)
    return "TIMEOUT", 0


def search_hit(auth: str, dataset_id: str, keyword: str) -> bool:
    r = requests.post(
        f"{BASE}/api/v1/datasets/{dataset_id}/search",
        headers=headers(auth),
        json={"question": keyword},
        timeout=60,
    )
    body = r.json()
    if body.get("code") != 0:
        return False
    data = body.get("data") or {}
    chunks = data.get("chunks") or data.get("documents") or []
    if not isinstance(chunks, list):
        return False
    hay = json.dumps(chunks, ensure_ascii=False).lower()
    return keyword.lower() in hay or len(chunks) > 0


def run_smoke() -> dict:
    auth = login()
    ds_name = f"TBOX-G1-SMOKE-{DATE_TAG}"
    dataset_id = create_dataset(auth, ds_name)
    sample_dir = Path(os.environ.get("TBOX_SMOKE_TMP", "/tmp/tbox-g1-smoke")) / DATE_TAG
    samples = make_samples(sample_dir)
    results: list[FormatResult] = []

    for label, (path, kw) in samples.items():
        note_parts: list[str] = []
        doc_id = upload_file(auth, dataset_id, path)
        upload_ok = bool(doc_id)
        if not upload_ok:
            results.append(FormatResult(label, kw, False, False, False, "—", 0, "upload failed"))
            continue

        if not trigger_parse(auth, dataset_id, doc_id):
            note_parts.append("parse API failed")
        run, chunks = wait_parse(auth, dataset_id, doc_id)
        parse_ok = run in ("3", "DONE") and chunks > 0
        if run in ("4", "FAIL"):
            note_parts.append("parse FAIL")
        if run == "TIMEOUT":
            note_parts.append(f"parse timeout>{PARSE_TIMEOUT_SEC}s")

        hit = False
        if parse_ok:
            hit = search_hit(auth, dataset_id, kw)
            if not hit:
                note_parts.append("search no hit")

        results.append(
            FormatResult(
                label,
                kw,
                upload_ok,
                parse_ok,
                hit,
                run,
                chunks,
                "; ".join(note_parts) if note_parts else "ok",
            )
        )

    return {
        "base_url": BASE,
        "dataset_id": dataset_id,
        "dataset_name": ds_name,
        "email": EMAIL,
        "date": DATE_TAG,
        "results": [asdict(r) for r in results],
    }


if __name__ == "__main__":
    print(json.dumps(run_smoke(), ensure_ascii=False, indent=2))
